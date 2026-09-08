from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/research/report-source.md"
OUTPUT = ROOT / "docs/research/Tally-competitive-feature-research.docx"

NAVY = "17365D"
BLUE = "DDEBF7"
PALE = "F5F8FC"
GRAY = "D9D9D9"
TEXT = RGBColor(31, 41, 55)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "3155D9")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text):
    pattern = re.compile(r"(https?://\S+|`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("http"):
            url = token.rstrip(".,;)")
            suffix = token[len(url):]
            add_hyperlink(paragraph, url, url)
            if suffix:
                paragraph.add_run(suffix)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Aptos Mono"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(68, 68, 68)
        else:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def configure(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, before, after in (("Title", 25, 0, 10), ("Heading 1", 17, 15, 7), ("Heading 2", 13, 11, 5), ("Heading 3", 11, 8, 4)):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "TALLY  |  PRODUCT RESEARCH"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.bold = True
            run.font.color.rgb = RGBColor(90, 102, 116)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("8 September 2026")
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.45, 1.75, 2.15, 1.7] if len(rows[0]) == 4 else [7.0 / len(rows[0])] * len(rows[0])
    twips = [round(w * 1440) for w in widths]
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(twips)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for i, row in enumerate(rows):
        tr_pr = table.rows[i]._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(twips[j]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if i == 0:
                set_cell_fill(cell, NAVY)
            elif i % 2 == 0:
                set_cell_fill(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_inline(p, value)
            for run in p.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8.3 if len(rows[0]) == 4 else 9)
                if i == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
    # repeat header
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc)

    # Cover
    title = doc.add_paragraph(style="Title")
    title.add_run("Tally Competitive Feature Research")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Current strengths, market gaps and a practical product roadmap").bold = True
    subtitle.paragraph_format.space_after = Pt(20)
    meta = doc.add_paragraph("Prepared for the Tally product owner\n8 September 2026\nIreland and UK consumer personal finance")
    meta.paragraph_format.space_after = Pt(20)
    intro = doc.add_paragraph()
    intro.add_run("Main finding  ").bold = True
    intro.add_run("Tally's strongest position is a shared, auditable household money ledger. The best next move is a transparent cash-flow forecast, followed by editable automation rules and more adaptable budgets.")
    doc.add_page_break()

    in_table = False
    table_rows = []
    skipped_front = 0
    for line in lines:
        if skipped_front < 5:
            if line.startswith("# ") or line.startswith("Audience:") or line.startswith("Date:") or line.startswith("Scope:") or not line.strip():
                skipped_front += 1
                continue
        if line.startswith("| "):
            parts = [x.strip() for x in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r"[-: ]+", x) for x in parts):
                continue
            table_rows.append(parts)
            in_table = True
            continue
        if in_table:
            add_table(doc, table_rows)
            table_rows = []
            in_table = False
        if not line.strip():
            continue
        if line.startswith("## "):
            text = line[3:]
            if text in {"Competitor profiles", "Recommended roadmap", "Claim to source ledger"}:
                doc.add_page_break()
            doc.add_heading(text, level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line[2:])
        elif line.startswith("# "):
            continue
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
    if table_rows:
        add_table(doc, table_rows)

    # Add page numbers in footers using a field.
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.add_run("  |  ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)

    doc.core_properties.title = "Tally Competitive Feature Research"
    doc.core_properties.subject = "Personal finance competitor features and Tally roadmap"
    doc.core_properties.author = "OpenAI Codex"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
